#!/usr/bin/env python3
"""Build the 7-slide deck on superposition, entanglement, and superposition tic-tac-toe.

Slides 2 to 4 cover the physics: superposition, qubits in quantum computers,
and entanglement. Slides 5 to 7 teach the game in
Superposition_Only_TicTacToe_Offline.html. The game diagrams reuse the game's
look: a red X, a blue O, and dashed spooky marks labeled with the move number
and half (A or B).

This script is the source of truth. Edit the text here and rebuild:

    python3 build-superposition-slides.py
"""
import argparse
import io
import math
from datetime import datetime, timezone
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FOLDER = Path(__file__).resolve().parent
OUTPUT_NAME = 'Superposition_Only_TicTacToe_Slides.pptx'
TOTAL = 7

# Colors from Superposition_Only_TicTacToe_Offline.html.
RED, BLUE, BOARD = 'B41524', '154C9C', 'BA1828'
INK, GRAY, NUMBER = '202435', '535969', '626575'
PANEL, EDGE, ARROW, WHITE = 'F0F3F8', 'D5D8DE', 'C9CED8', 'FFFFFF'
PAIR = {1: '975000', 2: '086D75', 3: '673BB7', 5: '316226'}   # the game's color for each move
METAL, CRYSTAL = '8A90A0', 'DCE7F7'                             # physics diagrams

FOOTER = 'GHOST GRID  ·  Superposition Tic-Tac-Toe'
IMAGE_ALT = ('Ghost Grid: Quantum Tic-Tac-Toe. A smiling ghost floats across a red '
             'tic-tac-toe board with solid and dashed X and O marks.')


# ---------------------------------------------------------------- helpers
def add_run(paragraph, words, *, size, color, bold=False):
    run = paragraph.add_run()
    run.text = words
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def list_marker(paragraph, kind, color=RED, indent=0.42):
    props = paragraph._p.get_or_add_pPr()
    props.set('marL', str(Inches(indent)))
    props.set('indent', str(-Inches(indent)))
    etree.SubElement(etree.SubElement(props, qn('a:buClr')), qn('a:srgbClr')).set('val', color)
    etree.SubElement(props, qn('a:buFont')).set('typeface', 'Arial')
    if kind == 'number':
        etree.SubElement(props, qn('a:buAutoNum')).set('type', 'arabicPeriod')
    else:
        etree.SubElement(props, qn('a:buChar')).set('char', '•')


def text(slide, x, y, w, h, lines, *, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, after=0, marker=None):
    """Text box. Each line is a string or a list of (words, overrides) runs."""
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(after)
        for words, style in (line if isinstance(line, list) else [(line, {})]):
            add_run(paragraph, words, size=style.get('size', size),
                    color=style.get('color', color), bold=style.get('bold', bold))
        if marker:
            list_marker(paragraph, marker)
    return frame


def shape(slide, x, y, w, h, *, fill=None, line=None, width=0, dash=False,
          kind=MSO_SHAPE.RECTANGLE, radius=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.shadow.inherit = False
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        item.fill.background()
    if line:
        item.line.color.rgb = RGBColor.from_string(line)
        item.line.width = Pt(width)
        if dash:
            item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        item.line.fill.background()
    if radius is not None:
        item.adjustments[0] = radius
    return item


def square(slide, x, y, size, number=None):
    shape(slide, x, y, size, size, fill=WHITE, line=BOARD, width=3)
    if number is not None:
        text(slide, x + 0.08, y + 0.05, 0.4, 0.3, [str(number)], size=min(13, round(size * 12)),
             bold=True, color=NUMBER)


def real(slide, x, y, size, player):
    text(slide, x, y, size, size, [player], size=round(size * 50), bold=True,
         color=RED if player == 'X' else BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def spooky(slide, x, y, size, player, move, half, numbered=False):
    """Dashed spooky mark in the square at (x, y), labeled like the game.

    In a numbered square the mark sits lower so the square number stays clear.
    """
    if numbered:
        w, h, top, big, small = size * 0.6, size * 0.62, y + size * 0.3, size * 26, size * 9.5
    else:
        w, h, top, big, small = size * 0.66, size * 0.7, y + size * 0.2, size * 30, size * 10
    mark = shape(slide, x + (size - w) / 2, top, w, h,
                 line=PAIR[move], width=2.5, dash=True, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    frame = mark.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = frame.paragraphs[0]
    first.alignment = PP_ALIGN.CENTER
    add_run(first, player, size=round(big), color=PAIR[move], bold=True)
    second = frame.add_paragraph()
    second.alignment = PP_ALIGN.CENTER
    add_run(second, f'{move} · {half}', size=round(small), color=PAIR[move], bold=True)


def die(slide, x, y, size, face):
    shape(slide, x, y, size, size, fill=WHITE, line=INK, width=3,
          kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    spots = {4: [(0.28, 0.28), (0.72, 0.28), (0.28, 0.72), (0.72, 0.72)]}
    dot = size * 0.17
    for fx, fy in spots[face]:
        shape(slide, x + fx * size - dot / 2, y + fy * size - dot / 2, dot, dot,
              fill=INK, kind=MSO_SHAPE.OVAL)


def arrow(slide, x, y, w, h):
    shape(slide, x, y, w, h, fill=ARROW, kind=MSO_SHAPE.DOWN_ARROW)


def label(slide, x, y, w, words, align=PP_ALIGN.CENTER):
    text(slide, x, y, w, 0.42, [words], size=18, bold=True, color=GRAY, align=align)


def band(slide, x, y, w, h):
    """Light panel with a red edge, like the game's status line."""
    shape(slide, x, y, w, h, fill=PANEL)
    shape(slide, x, y, 0.09, h, fill=RED)


def card(slide, x, y, w, h):
    shape(slide, x, y, w, h, fill=WHITE, line=EDGE, width=1.5,
          kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)


def badge(slide, x, y, w, h, words, *, fill=WHITE, line=None, color=INK, size=18,
          kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None):
    """Shape with centered bold text."""
    item = shape(slide, x, y, w, h, fill=fill, line=line, width=2, kind=kind, radius=radius)
    frame = item.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    add_run(paragraph, words, size=size, color=color, bold=True)
    return item


def arrowhead(item):
    tail = etree.SubElement(item.line._get_or_add_ln(), qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')


def beam(slide, x1, y1, x2, y2, *, color=RED, width=3):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    item.line.color.rgb = RGBColor.from_string(color)
    item.line.width = Pt(width)
    return item


def wave(slide, x, y, w, h, *, color=RED, cycles=2.5, steps=48):
    """Wavy line with an arrowhead, for a photon."""
    points = [(Inches(x + w * i / steps),
               Inches(y + h / 2 - h / 2 * math.sin(2 * math.pi * cycles * i / steps)))
              for i in range(steps + 1)]
    builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=1.0)
    builder.add_line_segments(points[1:], close=False)
    item = builder.convert_to_shape()
    item.shadow.inherit = False
    item.fill.background()
    item.line.color.rgb = RGBColor.from_string(color)
    item.line.width = Pt(3)
    arrowhead(item)
    return item


def icon_ion(slide, x, y, s):
    """Charged atom held between two trap electrodes and hit by a laser."""
    for top in (0.1, 0.8):
        shape(slide, x + 0.12 * s, y + top * s, 0.76 * s, 0.1 * s, fill=METAL,
              kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    beam(slide, x, y + 0.5 * s, x + 0.36 * s, y + 0.5 * s)
    badge(slide, x + 0.36 * s, y + 0.33 * s, 0.34 * s, 0.34 * s, '+', fill=BLUE, color=WHITE,
          size=16, kind=MSO_SHAPE.OVAL)


def icon_circuit(slide, x, y, s):
    """Superconducting loop broken by a junction."""
    shape(slide, x + 0.14 * s, y + 0.14 * s, 0.72 * s, 0.72 * s, line=INK, width=4,
          kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    shape(slide, x + 0.8 * s, y + 0.4 * s, 0.12 * s, 0.2 * s, fill=WHITE)
    text(slide, x + 0.68 * s, y + 0.34 * s, 0.36 * s, 0.32 * s, ['×'], size=22, bold=True,
         color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def icon_photon(slide, x, y, s):
    wave(slide, x + 0.02 * s, y + 0.3 * s, 0.9 * s, 0.4 * s)


def qubit_card(slide, y, title, words, icon):
    x, w, h = 7.3, 5.45, 1.45
    card(slide, x, y, w, h)
    icon(slide, x + 0.18, y + 0.2, 1.05)
    text(slide, x + 1.45, y + 0.15, w - 1.6, 0.42, [title], size=20, bold=True)
    text(slide, x + 1.45, y + 0.6, w - 1.6, 0.8, [words], size=16, color=GRAY)


def notes(slide, words):
    slide.notes_slide.notes_text_frame.text = words


def content_slide(prs, number, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.6, 0.35, 12.1, 0.9, [title], size=34, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    shape(slide, 0.6, 1.36, 1.2, 0.07, fill=RED)
    text(slide, 0.6, 6.98, 8.0, 0.32, [FOOTER], size=12, color=GRAY)
    text(slide, 11.2, 6.98, 1.55, 0.32, [f'{number} / {TOTAL}'], size=12, color=GRAY,
         align=PP_ALIGN.RIGHT)
    return slide


def ghost_png(path, pad=16):
    """Trim the white border from the ghost image so it shows larger."""
    with Image.open(path) as source:
        image = source.convert('RGB')
    ink = image.convert('L').point(lambda v: 255 if v < 230 else 0)
    left, top, right, bottom = ink.getbbox()
    box = (max(left - pad, 0), max(top - pad, 0),
           min(right + pad, image.width), min(bottom + pad, image.height))
    stream = io.BytesIO()
    image.crop(box).save(stream, format='PNG')
    stream.seek(0)
    return stream


# ----------------------------------------------------------------- slides
def slide_title(prs, image_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.8, 1.6, 6.4, 2.4, ['Superposition', 'Tic-Tac-Toe'], size=60, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    shape(slide, 0.8, 4.25, 1.6, 0.09, fill=RED)
    text(slide, 0.8, 4.55, 6.3, 0.6, ['A quantum game for two players'], size=26, color=GRAY)
    text(slide, 0.8, 6.45, 6.3, 0.4,
         ['Adapted from Allan Goff, American Journal of Physics 74 (2006)'], size=12, color=GRAY)
    picture = slide.shapes.add_picture(ghost_png(image_path), Inches(7.2), Inches(0.75),
                                       height=Inches(6.0))
    picture.name = 'Ghost Grid image'
    picture._element.nvPicPr.cNvPr.set('descr', IMAGE_ALT)
    notes(slide, 'Before class, open Superposition_Only_TicTacToe_Offline.html in a browser. '
                 'The game runs without internet. Click Full screen to project it.')


def slide_superposition(prs):
    slide = content_slide(prs, 2, 'Superposition: two places until measured')
    text(slide, 0.6, 1.75, 6.6, 4.9, [
        'Tiny quantum objects, like electrons, follow their own rules.',
        'Before a measurement, an electron can seem to be in two places at once.',
        'Measuring finds it in one place.',
        'No one can predict which place.',
    ], size=24, after=22, marker='dot')
    label(slide, 7.4, 1.6, 5.4, 'Before measuring')
    for x, half in ((8.3, 'A'), (10.4, 'B')):
        square(slide, x, 2.1, 1.5)
        spooky(slide, x, 2.1, 1.5, 'X', 1, half)
    arrow(slide, 9.8, 3.72, 0.6, 0.72)
    text(slide, 10.55, 3.84, 2.2, 0.5, ['Measure: roll a die'], size=18, bold=True, color=RED)
    label(slide, 7.4, 4.52, 5.4, 'After measuring')
    square(slide, 8.3, 5.02, 1.5)
    square(slide, 10.4, 5.02, 1.5)
    real(slide, 10.4, 5.02, 1.5, 'X')
    notes(slide, 'In the game, a die stands in for the random result of a measurement.')


def slide_qubits(prs):
    slide = content_slide(prs, 3, 'Quantum computers use qubits')
    text(slide, 0.6, 1.7, 6.4, 5.0, [
        'A regular computer bit is either 0 or 1.',
        'A quantum bit, or qubit, can be in a superposition of 0 and 1.',
        'A precise laser or microwave pulse puts a qubit into superposition.',
        'Heat and noise ruin a superposition, so qubits need shielding.',
        'With enough qubits, some problems could be solved much faster.',
    ], size=22, after=14, marker='dot')
    label(slide, 7.3, 1.6, 5.45, 'Three ways to build a qubit')
    qubit_card(slide, 2.05, 'Trapped ion',
               'One charged atom floats in a trap. Lasers control it.', icon_ion)
    qubit_card(slide, 3.6, 'Superconducting circuit',
               'A tiny circuit is kept colder than outer space. Microwaves control it.', icon_circuit)
    qubit_card(slide, 5.15, 'Photon',
               'A photon is a particle of light. The way it wiggles stores 0 or 1.', icon_photon)
    notes(slide, 'Google and IBM build qubits from superconducting circuits. '
                 'IonQ and Quantinuum use trapped ions. '
                 'A common myth says a quantum computer tries every answer at once. It does not. '
                 'It uses superposition and interference to make right answers more likely.')


def slide_entanglement(prs):
    slide = content_slide(prs, 4, 'Entanglement links two particles')
    text(slide, 0.6, 1.7, 6.4, 5.0, [
        'Particles become entangled when they interact or are created together.',
        'Entangled particles share one quantum state, even far apart.',
        'In a simple pair, if one shows 0, the other shows 0 too.',
        'Each result is random, so entanglement cannot send messages faster than light.',
        'Quantum computers entangle qubits so they can work together.',
    ], size=22, after=14, marker='dot')
    label(slide, 7.3, 1.6, 5.45, 'How scientists make entangled photons')
    badge(slide, 7.45, 2.75, 1.15, 0.62, 'Laser', fill=INK, color=WHITE, size=16, radius=0.15)
    beam(slide, 8.6, 3.06, 9.55, 3.06, width=4)
    shape(slide, 9.55, 2.61, 0.9, 0.9, fill=CRYSTAL, line=BLUE, width=2, kind=MSO_SHAPE.DIAMOND)
    text(slide, 9.3, 3.55, 1.4, 0.35, ['Crystal'], size=14, color=GRAY, align=PP_ALIGN.CENTER)
    for top, name in ((1.98, 'A'), (3.5, 'B')):
        beam(slide, 10.45, 3.06, 11.95, top + 0.32)
        badge(slide, 11.95, top, 0.64, 0.64, name, fill=BLUE, color=WHITE, kind=MSO_SHAPE.OVAL)
    text(slide, 7.3, 4.3, 5.45, 0.4, ['One photon can split into two entangled photons.'],
         size=16, color=GRAY, align=PP_ALIGN.CENTER)
    label(slide, 7.3, 4.95, 5.45, 'Measure both photons')
    for x, words in ((7.55, 'A: 0'), (8.7, 'B: 0'), (10.4, 'A: 1'), (11.55, 'B: 1')):
        badge(slide, x, 5.4, 1.05, 0.55, words, line=BLUE, color=BLUE, radius=0.25)
    text(slide, 9.75, 5.45, 0.65, 0.45, ['or'], size=16, color=GRAY, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 7.3, 6.07, 5.45, 0.4, ['The results match. Which one you get is random.'],
         size=16, color=GRAY, align=PP_ALIGN.CENTER)
    notes(slide, 'Einstein doubted entanglement and called it “spooky action at a distance.” '
                 'Goff named the game’s spooky marks after that phrase. '
                 'The 2022 Nobel Prize in Physics honored experiments with entangled photons. '
                 'It went to Alain Aspect, John Clauser, and Anton Zeilinger. '
                 'In the game, two pairs that share a square act a little like an entangled pair.')


def slide_moves(prs):
    slide = content_slide(prs, 5, 'On your turn, pick one kind of move')
    card(slide, 0.6, 1.65, 5.9, 3.45)
    text(slide, 0.95, 1.85, 5.2, 0.5, ['Classical move'], size=26, bold=True)
    text(slide, 0.95, 2.4, 5.2, 0.9,
         ['Put one real X or O in a square that is completely empty.'], size=22)
    square(slide, 2.85, 3.5, 1.4, 5)
    real(slide, 2.85, 3.5, 1.4, 'X')
    card(slide, 6.85, 1.65, 5.9, 3.45)
    text(slide, 7.2, 1.85, 5.2, 0.5, ['Superposition move'], size=26, bold=True)
    text(slide, 7.2, 2.4, 5.2, 0.9,
         ['Put a matching pair of spooky marks in two squares, A and B.'], size=22)
    for x, number, half in ((8.25, 3, 'A'), (9.95, 7, 'B')):
        square(slide, x, 3.5, 1.4, number)
        spooky(slide, x, 3.5, 1.4, 'O', 2, half, numbered=True)
    band(slide, 0.6, 5.3, 12.15, 1.45)
    text(slide, 0.98, 5.45, 11.5, 1.2, [
        'A square can hold up to two spooky marks.',
        'A real mark locks its square.',
        'Placing stops when every square has at least one mark.',
    ], size=20, after=4, marker='dot')
    notes(slide, 'On screen, choose Classical or Superposition above the playing pad. '
                 'Then click a square or press a key from 1 to 9. '
                 'Each pair gets its own color and move number.')


def slide_measure(prs):
    slide = content_slide(prs, 6, 'Measure your own pairs with a die')
    text(slide, 0.6, 1.7, 6.7, 3.7, [
        'Measure when every square has a mark.',
        'Take turns. Measure only your own pairs.',
        'Pick one of your pairs and roll the die.',
        'Odd keeps A. Even keeps B.',
        'The kept mark becomes real. The other mark disappears.',
    ], size=24, after=10, marker='number')
    label(slide, 7.55, 2.02, 1.1, 'Before', align=PP_ALIGN.LEFT)
    for x, number, half in ((8.75, 1, 'A'), (10.15, 5, 'B')):
        square(slide, x, 1.65, 1.25, number)
        spooky(slide, x, 1.65, 1.25, 'X', 1, half, numbered=True)
    die(slide, 11.78, 1.78, 0.95, 4)
    text(slide, 11.55, 2.8, 1.4, 0.35, ['Roll: 4'], size=16, bold=True, align=PP_ALIGN.CENTER)
    arrow(slide, 9.83, 3.0, 0.5, 0.55)
    label(slide, 7.55, 4.1, 1.1, 'After', align=PP_ALIGN.LEFT)
    square(slide, 8.75, 3.65, 1.25, 1)
    square(slide, 10.15, 3.65, 1.25, 5)
    real(slide, 10.15, 3.65, 1.25, 'X')
    text(slide, 8.2, 4.97, 4.5, 0.4, ['4 is even, so B stays in square 5.'], size=16,
         color=GRAY, align=PP_ALIGN.CENTER)
    band(slide, 0.6, 5.5, 12.15, 1.3)
    text(slide, 0.98, 5.63, 11.5, 1.05, [[
        ('Shared squares: ', {'bold': True}),
        ('two pairs cannot both stay in one square. If a kept mark lands there, '
         'the other pair becomes real in its other square.', {}),
    ]], size=20)
    notes(slide, 'On screen, click one of your spooky marks, then click Roll die. '
                 'The game makes any forced moves and lists them after the roll. '
                 'If a player has no pairs left, the other player keeps measuring. '
                 'Squares that empty out open again for the next round.')


def slide_win(prs):
    slide = content_slide(prs, 7, 'Win after all pairs are measured')
    text(slide, 0.6, 1.7, 7.4, 3.4, [
        'Three real marks in a row wins.',
        'Spooky marks do not count.',
        'A row for both players is a draw.',
        'A full board with no row is a draw.',
        'Otherwise, place and measure again.',
    ], size=24, after=10, marker='dot')
    size, left, top = 0.95, 9.15, 1.5
    for index in range(9):
        square(slide, left + (index % 3) * size, top + (index // 3) * size, size, index + 1)
    for column, (move, half) in enumerate(((1, 'A'), (3, 'B'), (5, 'A'))):
        spooky(slide, left + column * size, top, size, 'X', move, half, numbered=True)
    text(slide, 8.4, 4.45, 4.35, 0.7, ['Not a win yet.', 'All three X marks are spooky.'],
         size=18, color=GRAY, align=PP_ALIGN.CENTER)
    band(slide, 0.6, 5.35, 12.15, 1.45)
    text(slide, 0.98, 5.47, 11.5, 0.4, ['Talk about it'], size=20, bold=True)
    text(slide, 0.98, 5.9, 11.5, 0.85, [
        'Why do three spooky marks in a row not count as a win?',
        'Two pairs share a square. What does measuring one tell you about the other?',
    ], size=18, after=2, marker='dot')
    notes(slide, 'Spooky marks are not real until they are measured. '
                 'That is why a row of spooky marks does not win. '
                 'Measurement decides which marks become real.')


def properties(prs):
    core = prs.core_properties
    core.title = 'Superposition Tic-Tac-Toe'
    core.subject = ('Seven slides on superposition, qubits, entanglement, and the Ghost Grid '
                    'superposition tic-tac-toe game')
    core.keywords = ('quantum tic-tac-toe, superposition, entanglement, quantum computing, '
                     'grade 6, Ghost Grid')
    core.author = core.last_modified_by = core.comments = ''
    core.created = core.modified = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    core.revision = 1


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--image', type=Path, default=FOLDER / 'GhostGrid.png')
    parser.add_argument('--out', type=Path, default=FOLDER / OUTPUT_NAME)
    args = parser.parse_args()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    slide_title(prs, args.image)
    slide_superposition(prs)
    slide_qubits(prs)
    slide_entanglement(prs)
    slide_moves(prs)
    slide_measure(prs)
    slide_win(prs)
    properties(prs)
    prs.save(args.out)
    print('Built:', args.out)


if __name__ == '__main__':
    main()
