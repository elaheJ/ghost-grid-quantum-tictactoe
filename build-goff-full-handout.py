#!/usr/bin/env python3
"""Build the Grade 6 handout for Goff's full quantum tic-tac-toe.

The full game combines superposition and entanglement. It is the game in
Goff_TicTacToe_Offline.html. The rules follow Allan Goff, American Journal of
Physics 74 (2006), pp. 962-965.

This script is the source of truth for the handout. Edit the text here and
rebuild. Do not hand-edit the .docx.

    python3 build-goff-full-handout.py

The page follows the older Codex-built handout's layout: US Letter, rules
on top, a large numbered board with doubled lines below, and the citation in
the footer. A band at the top adds GhostGrid.png and three big ideas.
"""
import argparse
import io
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from PIL import Image

FOLDER = Path(__file__).resolve().parent
OUTPUT_NAME = 'Quantum_TicTacToe_Goff_Grade6_Handout.docx'

# ------------------------------------------------------------------ text
TITLE = 'SUPERPOSITION + ENTANGLEMENT'
SUBTITLE = 'Goff’s full game  •  Two players  •  Pencils and erasers'

IDEAS = [
    ('Superposition:', 'a spooky pair is one move in two places at once.'),
    ('Entanglement:', 'pairs that share a square are linked, so they affect each other.'),
    ('Collapse:', 'a measurement turns a spooky pair into one real mark.'),
]

# X_1 prints as X with a subscript 1.
RULES = [
    ('Spooky pairs (superposition).',
     'X goes first. On each move, write your letter in two different squares. '
     'Label both with the move number: X_1 and X_1, then O_2 and O_2, then X_3 and X_3.'),
    ('Shared squares (entanglement).',
     'Marks from different moves may share a square, so write small. '
     'Two moves that share a square cannot both end up there. '
     'Do not write in a square with a real mark.'),
    ('Loops.',
     'A move closes a loop if a chain of pairs already joins its two squares. '
     'Example: X_1 is in squares 1 and 2. O_2 is in 2 and 5. '
     'X_3 in 1 and 5 closes a loop. Underline the linked marks.'),
    ('Collapse (measurement).',
     'The player who did NOT close the loop chooses which square the newest move keeps. '
     'No dice are used. In the example, O picks square 1 or square 5 for X_3.'),
    ('Chain reaction.',
     'Write the chosen mark large. It is real now. Cross out its twin. '
     'Cross out other spooky marks in that square and make their twins real. '
     'Repeat for each new real mark. Pairs not linked to the loop stay spooky. '
     'If O picks square 1, X_3 is real in 1, X_1 in 2, and O_2 in 5.'),
    ('Win or keep going.',
     'Three real marks in a row wins: across, down, or diagonal. Spooky marks do not count. '
     'Check for a win after each collapse. If no one has won, the chooser makes the next move.'),
    ('Several rows at once.',
     'Each row’s finish number is its largest move number. '
     'Rows with the lowest finish number score 1 point each. Later rows score ½ point each. '
     'The higher total wins.'),
    ('Game end.',
     'The game ends after move 9 and any collapse. With no real winning row, it is a tie. '
     'If only one square is open at move 9, X writes both marks there. They make one real X.'),
]

FOOTER = 'Rules: Allan Goff, American Journal of Physics 74 (2006), pp. 962–965. DOI: 10.1119/1.2213635'
IMAGE_ALT = ('Ghost Grid: Quantum Tic-Tac-Toe. A smiling ghost floats across a red '
             'tic-tac-toe board with solid and dashed X and O marks.')

# ---------------------------------------------------------------- layout
# Sizes are in twips (1 inch = 1440) unless the name says otherwise.
RED = 'D3061C'      # sampled from GhostGrid.png
GRAY = '555555'
INK = '202020'
PAGE_W, PAGE_H = 12240, 15840
MARGIN_X, MARGIN_Y = 720, 576
TEXT_W = PAGE_W - 2 * MARGIN_X
BODY_PT, BODY_LINE = 11, 246        # same as the first handout
IDEA_PT, IDEA_LINE = 10.5, 236
IMAGE_IN = 1.5
IMAGE_COL = 2300
ROW_H = 2290                        # same as the first handout
BOARD_GAP = 200


def el(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(f'w:{key}'), str(value))
    return node


def run(paragraph, text, *, size=None, bold=False, color=None, subscript=False):
    r = paragraph.add_run(text)
    if bold:
        r.bold = True
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if subscript:
        r.font.subscript = True
    return r


def rich(paragraph, text, **font):
    """Write text, printing X_1 as X with a subscript 1."""
    for piece in re.split(r'([XO]_\d)', text):
        if re.fullmatch(r'[XO]_\d', piece):
            run(paragraph, piece[0], **font)
            run(paragraph, piece[2], subscript=True, **font)
        elif piece:
            run(paragraph, piece, **font)


def fmt(paragraph, *, style=None, before=0, after=0, line=None, keep=False, align=None,
        left=None, hanging=None):
    if style:
        paragraph.style = style
    pf = paragraph.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    if line:
        pf.line_spacing = Twips(line)
    if keep:
        pf.keep_together = True
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Twips(left)
    if hanging is not None:
        pf.first_line_indent = Twips(-hanging)
    return paragraph


def left_rule(paragraph, color=RED):
    """Thin colored bar on the left of a paragraph."""
    border = el('w:pBdr')
    border.append(el('w:left', val='single', sz=18, space=6, color=color))
    paragraph._p.get_or_add_pPr().insert_element_before(
        border, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
        'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi',
        'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing',
        'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
        'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
        'w:pPrChange')


def table_props(table, widths, borders, margins):
    """Fixed-width centered table, with children written in schema order."""
    tbl = table._tbl
    props = tbl.tblPr
    for child in list(props):
        props.remove(child)
    props.append(el('w:tblW', w=sum(widths), type='dxa'))
    props.append(el('w:jc', val='center'))
    edges = el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if side in borders:
            val, size = borders[side]
            edges.append(el(f'w:{side}', val=val, sz=size, space=0, color=INK))
        else:
            edges.append(el(f'w:{side}', val='nil'))
    props.append(edges)
    props.append(el('w:tblLayout', type='fixed'))
    cell_margins = el('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        cell_margins.append(el(f'w:{side}', w=margins[side], type='dxa'))
    props.append(cell_margins)
    props.append(el('w:tblLook', val='0000', firstRow=0, lastRow=0, firstColumn=0,
                    lastColumn=0, noHBand=1, noVBand=1))
    for column, width in zip(tbl.tblGrid.findall(qn('w:gridCol')), widths):
        column.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)


def ghost_png(path, pad=16):
    """Trim the white border from the ghost image so it prints larger."""
    with Image.open(path) as source:
        image = source.convert('RGB')
    ink = image.convert('L').point(lambda v: 255 if v < 230 else 0)
    left, top, right, bottom = ink.getbbox()
    box = (max(left - pad, 0), max(top - pad, 0),
           min(right + pad, image.width), min(bottom + pad, image.height))
    stream = io.BytesIO()
    image.crop(box).save(stream, format='PNG')
    stream.seek(0)
    return stream


def setup(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(PAGE_W), Twips(PAGE_H)
    section.left_margin = section.right_margin = Twips(MARGIN_X)
    section.top_margin = section.bottom_margin = Twips(MARGIN_Y)
    section.header_distance, section.footer_distance = Twips(216), Twips(230)

    defaults = doc.styles.element.find(qn('w:docDefaults'))
    fonts = defaults.find(qn('w:rPrDefault')).find(qn('w:rPr')).find(qn('w:rFonts'))
    for attr in ('asciiTheme', 'hAnsiTheme'):
        fonts.attrib.pop(qn(f'w:{attr}'), None)
    fonts.set(qn('w:ascii'), 'Arial')
    fonts.set(qn('w:hAnsi'), 'Arial')

    for name in ('Normal', 'Body Text'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(BODY_PT)
        style.paragraph_format.space_before = Twips(0)
        style.paragraph_format.space_after = Twips(40)
        style.paragraph_format.line_spacing = Twips(BODY_LINE)
    doc.styles['Body Text'].base_style = doc.styles['Normal']


def header_band(doc, image_path):
    table = doc.add_table(rows=1, cols=2)
    table_props(table, [IMAGE_COL, TEXT_W - IMAGE_COL], borders={},
                margins=dict(top=0, left=0, bottom=0, right=0))
    image_cell, text_cell = table.rows[0].cells
    image_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    picture = fmt(image_cell.paragraphs[0])
    picture.paragraph_format.line_spacing = 1.0     # exact spacing would clip the image
    shape = picture.add_run().add_picture(ghost_png(image_path), height=Inches(IMAGE_IN))
    shape._inline.docPr.set('name', 'Ghost Grid')
    shape._inline.docPr.set('descr', IMAGE_ALT)

    run(fmt(text_cell.paragraphs[0], after=30, line=380), TITLE, size=17, bold=True, color=RED)
    run(fmt(text_cell.add_paragraph(), after=90, line=240), SUBTITLE, size=10)
    for index, (term, text) in enumerate(IDEAS):
        idea = fmt(text_cell.add_paragraph(), style='Body Text', line=IDEA_LINE, left=170,
                   after=40 if index < len(IDEAS) - 1 else 0)
        left_rule(idea)
        run(idea, term + ' ', size=IDEA_PT, bold=True)
        run(idea, text, size=IDEA_PT)


def rules(doc):
    for number, (lead, text) in enumerate(RULES, start=1):
        rule = fmt(doc.add_paragraph(), style='Body Text', before=110 if number == 1 else 0,
                   after=40, line=BODY_LINE, keep=True, left=274, hanging=274)
        run(rule, f'{number}. {lead} ', bold=True)
        rich(rule, text)


def board(doc):
    run(fmt(doc.add_paragraph(), line=BOARD_GAP), '', size=1)
    table = doc.add_table(rows=3, cols=3)
    single, double = ('single', 12), ('double', 12)
    table_props(table, [TEXT_W // 3] * 3,
                borders=dict(top=single, left=single, bottom=single, right=single,
                             insideH=double, insideV=double),
                margins=dict(top=65, left=90, bottom=45, right=70))
    for r, row in enumerate(table.rows):
        row.height, row.height_rule = Twips(ROW_H), WD_ROW_HEIGHT_RULE.EXACTLY
        row._tr.get_or_add_trPr().insert(0, el('w:cantSplit'))
        for c, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            run(fmt(cell.paragraphs[0], line=240), str(3 * r + c + 1), size=10, color=GRAY)
    # Word needs a paragraph after a final table. Keep it tiny so it cannot spill.
    run(fmt(doc.add_paragraph(), line=20), '', size=1)


def footer(doc):
    note = doc.sections[0].footer.paragraphs[0]
    fmt(note, style='Footer', line=180, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(note, FOOTER, size=8, color=GRAY)


def properties(doc):
    core = doc.core_properties
    core.title = 'Quantum Tic-Tac-Toe: Superposition + Entanglement (Grade 6)'
    core.subject = 'One-page handout for Allan Goff’s full 2006 game, with the Ghost Grid image'
    core.keywords = 'quantum tic-tac-toe, superposition, entanglement, grade 6, Ghost Grid'
    core.author = core.last_modified_by = core.comments = ''
    core.created = core.modified = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    core.revision = 1


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--image', type=Path, default=FOLDER / 'GhostGrid.png')
    parser.add_argument('--out', type=Path, default=FOLDER / OUTPUT_NAME)
    args = parser.parse_args()

    doc = Document()
    body = doc.element.body
    for paragraph in body.findall(qn('w:p')):
        body.remove(paragraph)
    setup(doc)
    header_band(doc, args.image)
    rules(doc)
    board(doc)
    footer(doc)
    properties(doc)
    doc.save(args.out)
    print('Built:', args.out)


if __name__ == '__main__':
    main()
